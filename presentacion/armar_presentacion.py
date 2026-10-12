"""Genera presentacion/presentacion.pptx para la defensa oral del TP.

Estructura narrativa: problema -> metodo -> resultado aparente (100%) -> giro
(la baseline empata) -> hallazgo (la red mide altura, no tonalidad) -> leccion.
"""
import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Cm, Pt

BASE = os.path.dirname(os.path.abspath(__file__))
FIG = os.path.join(os.path.dirname(BASE), "informe", "figuras")
SALIDA = os.path.join(BASE, "presentacion.pptx")

AZUL = RGBColor(0x1F, 0x38, 0x64)
AZUL_CLARO = RGBColor(0x2E, 0x5C, 0x9A)
GRIS = RGBColor(0x44, 0x44, 0x44)
BLANCO = RGBColor(0xFF, 0xFF, 0xFF)
ROJO = RGBColor(0xC0, 0x39, 0x2B)
VERDE = RGBColor(0x1E, 0x7A, 0x46)
FONDO_CLARO = RGBColor(0xF4, 0xF7, 0xFB)

prs = Presentation()
prs.slide_width = Cm(33.87)
prs.slide_height = Cm(19.05)
AN, AL = prs.slide_width, prs.slide_height
VACIO = prs.slide_layouts[6]


def fondo(slide, color):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = color


def caja(slide, x, y, w, h, texto, size=18, bold=False, color=GRIS,
         align=PP_ALIGN.LEFT, fuente="Calibri", interlineado=1.18):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    lineas = texto.split("\n") if isinstance(texto, str) else texto
    for i, ln in enumerate(lineas):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = interlineado
        r = p.add_run()
        r.text = ln
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        r.font.name = fuente
    return tb


def vinetas(slide, x, y, w, h, items, size=17, color=GRIS, espacio=10):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        texto, destacado = (it if isinstance(it, tuple) else (it, False))
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(espacio)
        p.line_spacing = 1.15
        r = p.add_run()
        r.text = "•  " + texto
        r.font.size = Pt(size)
        r.font.bold = destacado
        r.font.color.rgb = AZUL if destacado else color
        r.font.name = "Calibri"
    return tb


def barra_titulo(slide, titulo, sub=None):
    band = slide.shapes.add_shape(1, 0, 0, AN, Cm(2.5))
    band.fill.solid()
    band.fill.fore_color.rgb = AZUL
    band.line.fill.background()
    band.shadow.inherit = False
    tf = band.text_frame
    tf.margin_left = Cm(1.1)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = titulo
    r.font.size = Pt(29)
    r.font.bold = True
    r.font.color.rgb = BLANCO
    r.font.name = "Calibri"
    if sub:
        caja(slide, Cm(1.1), Cm(2.75), AN - Cm(2.2), Cm(1.1), sub, size=16,
             color=AZUL_CLARO)


def pie_num(slide, n):
    caja(slide, AN - Cm(2.6), AL - Cm(1.4), Cm(1.8), Cm(0.9), str(n),
         size=12, color=GRIS, align=PP_ALIGN.RIGHT)


def tarjeta(slide, x, y, w, h, titulo, valor, color_valor, nota=None):
    sh = slide.shapes.add_shape(5, x, y, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = BLANCO
    sh.line.color.rgb = RGBColor(0xC9, 0xD4, 0xE4)
    sh.line.width = Pt(1.1)
    sh.shadow.inherit = False
    tf = sh.text_frame
    tf.word_wrap = True
    tf.margin_top = Cm(0.35)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = titulo
    r.font.size = Pt(15)
    r.font.color.rgb = GRIS
    r.font.name = "Calibri"
    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run()
    r2.text = valor
    r2.font.size = Pt(40)
    r2.font.bold = True
    r2.font.color.rgb = color_valor
    r2.font.name = "Calibri"
    if nota:
        p3 = tf.add_paragraph()
        p3.alignment = PP_ALIGN.CENTER
        r3 = p3.add_run()
        r3.text = nota
        r3.font.size = Pt(12.5)
        r3.font.color.rgb = GRIS
        r3.font.name = "Calibri"
    return sh


def tabla(slide, x, y, w, h, datos, ancho_cols=None, size=14):
    nf, nc = len(datos), len(datos[0])
    gf = slide.shapes.add_table(nf, nc, x, y, w, h)
    t = gf.table
    if ancho_cols:
        for i, cw in enumerate(ancho_cols):
            t.columns[i].width = cw
    for i, fila in enumerate(datos):
        for j, celda in enumerate(fila):
            texto, bold, col = (celda if isinstance(celda, tuple) else (celda, False, None))
            c = t.cell(i, j)
            c.text = ""
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.margin_left = Cm(0.25)
            c.margin_top = Cm(0.1)
            c.margin_bottom = Cm(0.1)
            p = c.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER if j > 0 else PP_ALIGN.LEFT
            r = p.add_run()
            r.text = texto
            r.font.size = Pt(size)
            r.font.name = "Calibri"
            r.font.bold = bold or i == 0
            if i == 0:
                r.font.color.rgb = BLANCO
                c.fill.solid()
                c.fill.fore_color.rgb = AZUL
            else:
                r.font.color.rgb = col or GRIS
                c.fill.solid()
                c.fill.fore_color.rgb = BLANCO if i % 2 else FONDO_CLARO
    return t


def imagen_centrada(slide, archivo, y, alto):
    ruta = os.path.join(FIG, archivo)
    if not os.path.exists(ruta):
        return
    from PIL import Image as PILImage
    iw, ih = PILImage.open(ruta).size
    h = alto
    w = int(h * iw / ih)
    if w > AN - Cm(4):
        w = AN - Cm(4)
        h = int(w * ih / iw)
    slide.shapes.add_picture(ruta, int((AN - w) / 2), y, width=w, height=h)


n = 0


def nueva(titulo=None, sub=None, claro=False):
    global n
    s = prs.slides.add_slide(VACIO)
    fondo(s, FONDO_CLARO if claro else BLANCO)
    if titulo:
        barra_titulo(s, titulo, sub)
    n += 1
    pie_num(s, n)
    return s


# ---------------------------------------------------------------- 1. Portada
s = nueva()
fondo(s, AZUL)
caja(s, Cm(2.5), Cm(5.2), AN - Cm(5), Cm(3), "Master of the Harmonica",
     size=58, bold=True, color=BLANCO, align=PP_ALIGN.CENTER)
caja(s, Cm(4), Cm(8.6), AN - Cm(8), Cm(2.6),
     "Clasificación de tonalidad en 24 clases con CNN 1D sobre audio crudo\n"
     "y derivación determinística de armónica blues",
     size=21, color=RGBColor(0xC9, 0xD9, 0xF0), align=PP_ALIGN.CENTER)
ln = s.shapes.add_shape(1, int(AN / 2 - Cm(4)), Cm(11.9), Cm(8), Pt(2.5))
ln.fill.solid()
ln.fill.fore_color.rgb = RGBColor(0x7F, 0xA8, 0xD8)
ln.line.fill.background()
ln.shadow.inherit = False
caja(s, Cm(4), Cm(13.1), AN - Cm(8), Cm(4),
     "TP Final  ·  Redes Neuronales y Deep Learning\n"
     "Maestría en Inteligencia Artificial  ·  Universidad de Palermo\n\n"
     "V. M. Danguise  ·  2026",
     size=17, color=BLANCO, align=PP_ALIGN.CENTER, interlineado=1.4)

# ---------------------------------------------------------------- 2. Problema
s = nueva("El problema", "Qué armónica agarro del estuche para tocar sobre este tema")
vinetas(s, Cm(1.6), Cm(4.3), Cm(16), Cm(12), [
    ("La armónica diatónica está construida en una tonalidad fija.", True),
    "Para improvisar hay que resolver dos cosas:",
    "   1. ¿En qué tonalidad está la canción?",
    "   2. ¿Qué armónica corresponde?",
    ("En blues NO se usa la armónica de la canción.", True),
    "Se toca en segunda posición (cross-harp): la armónica está\n    una cuarta justa por encima de la tónica.",
    "Motivo: así los bendings caen sobre las notas expresivas\n    de la escala de blues.",
], size=17)
bx = s.shapes.add_shape(5, Cm(19), Cm(5.5), Cm(13.2), Cm(8.2))
bx.fill.solid()
bx.fill.fore_color.rgb = BLANCO
bx.line.color.rgb = AZUL_CLARO
bx.line.width = Pt(1.6)
bx.shadow.inherit = False
caja(s, Cm(19.6), Cm(6.3), Cm(12), Cm(7),
     "Objetivo del sistema\n\n"
     "audio.mp3\n↓\n[ CNN 1D ]\n↓\n\"G maj\"  (conf. 0.92)\n↓\n[ regla +5 semitonos ]\n↓\n"
     "ARMÓNICA EN C  +  riff",
     size=17, bold=True, color=AZUL, align=PP_ALIGN.CENTER, interlineado=1.25)

# ---------------------------------------------------------------- 3. Diseño
s = nueva("Qué aprende la red y qué no", "Separar lo perceptual de lo determinístico")
tarjeta(s, Cm(1.8), Cm(4.6), Cm(14.5), Cm(5.4), "LO QUE APRENDE LA RED",
        "Tonalidad", AZUL_CLARO, "24 clases: 12 tónicas × {mayor, menor}")
tarjeta(s, Cm(17.6), Cm(4.6), Cm(14.5), Cm(5.4), "LO QUE NO APRENDE",
        "(tónica + 5) % 12", VERDE, "La armónica sale por regla musical exacta")
vinetas(s, Cm(1.8), Cm(11), Cm(30), Cm(7), [
    ("Entra: WAV/MP3 → mono → 16 kHz → ventanas de 4 s (64.000 muestras) de forma de onda CRUDA.", True),
    "No entra MIDI ni cromagrama: se quiere ver si la CNN aprende la estructura tonal desde la señal.",
    ("Sale: 24 logits, con confianza y Top-3. La armónica se deriva, no se predice.", True),
    "Implementar la regla como capa de red sería agregarle error a algo que ya es exacto.",
], size=16.5)

# ---------------------------------------------------------------- 4. Método
s = nueva("Método", "KeyCNN1D — 80.968 parámetros, entrenada desde cero")
tabla(s, Cm(1.6), Cm(4.4), Cm(18), Cm(9), [
    ["Bloque", "Operación", "Canales", "Kernel"],
    ["b1", "Conv1d (stride 8) + BN + ReLU + Pool", "1 → 16", "128"],
    ["b2", "Conv1d + BN + ReLU + Pool", "16 → 32", "16"],
    ["b3", "Conv1d + BN + ReLU + Pool", "32 → 64", "7"],
    ["b4", "Conv1d + BN + ReLU + Pool", "64 → 128", "5"],
    ["head", "AdaptiveAvgPool → 128 → 96 → 24", "—", "—"],
], ancho_cols=[Cm(2.4), Cm(9.4), Cm(3.4), Cm(2.8)], size=14)
vinetas(s, Cm(20.6), Cm(4.4), Cm(11.8), Cm(11), [
    ("Kernel inicial ancho (128 ≈ 8 ms)", True),
    "banco de filtros aprendido, no impuesto por FFT",
    ("AdaptiveAvgPool1d(1)", True),
    "la tonalidad es propiedad global del fragmento",
    ("Baseline obligatoria", True),
    "cromagrama CQT + perfiles Krumhansl-Schmuckler",
    ("Datos", True),
    "progresiones I-IV-V-I / i-iv-V-i sintéticas\n    1920 train / 480 val / 480 test",
], size=15.5, espacio=6)

# ---------------------------------------------------------------- 5. Resultado aparente
s = nueva("Resultado", "Test sintético, 480 ejemplos, 24 clases", claro=True)
tarjeta(s, Cm(6), Cm(5.4), Cm(9.5), Cm(6), "ACCURACY", "100 %", VERDE, "480 / 480")
tarjeta(s, Cm(18.4), Cm(5.4), Cm(9.5), Cm(6), "SCORE MIREX", "1,000", VERDE, "sin un solo error")
caja(s, Cm(3), Cm(12.4), AN - Cm(6), Cm(3),
     "Matriz de confusión 24×24 perfectamente diagonal.\n"
     "Convergencia a 99,6 % de validación en la PRIMERA época.",
     size=21, color=AZUL, align=PP_ALIGN.CENTER, interlineado=1.3)
caja(s, Cm(3), Cm(15.8), AN - Cm(6), Cm(2),
     "¿Terminamos? No.", size=25, bold=True, color=ROJO, align=PP_ALIGN.CENTER)

# ---------------------------------------------------------------- 6. El giro
s = nueva("El problema del 100 %", "Para eso servía la línea de base que pidió la cátedra")
tabla(s, Cm(5.5), Cm(4.8), Cm(23), Cm(4), [
    ["Método", "Accuracy", "MIREX"],
    [("CNN 1D (80.968 parámetros, 12 épocas)", False, None), ("1,000", True, VERDE), ("1,000", True, VERDE)],
    [("Cromagrama + Krumhansl-Schmuckler (0 parámetros)", False, None), ("1,000", True, ROJO), ("1,000", True, ROJO)],
], ancho_cols=[Cm(14), Cm(4.5), Cm(4.5)], size=16)
caja(s, Cm(2.5), Cm(10), AN - Cm(5), Cm(2.4),
     "Un método de los años 80, sin parámetros entrenables, empata con la red.",
     size=24, bold=True, color=ROJO, align=PP_ALIGN.CENTER)
vinetas(s, Cm(4.5), Cm(13.1), Cm(25), Cm(5), [
    ("El benchmark sintético está saturado: no discrimina entre métodos.", True),
    ("No hay evidencia de que la red haya aportado valor.", True),
    "Si el test no ordena métodos, hay que buscar una evaluación que sí lo haga.",
], size=18)

# ---------------------------------------------------------------- 7. OOD
s = nueva("Evaluación fuera de distribución", "11 escenarios · se perturba una dimensión por vez · la etiqueta no cambia")
imagen_centrada(s, "robustez.png", Cm(3.9), Cm(10.4))
caja(s, Cm(2), Cm(14.9), AN - Cm(4), Cm(3),
     "La CNN gana donde importa el contexto (progresión nueva: 92,5 % vs 50 %)…\n"
     "…y colapsa a 0,042 — el nivel del azar — ante un cambio de octava, donde la baseline queda en 1,000.",
     size=17, color=AZUL, align=PP_ALIGN.CENTER, interlineado=1.3)

# ---------------------------------------------------------------- 8. Tabla OOD
s = nueva("Los números", "Accuracy en 24 clases · azar = 1/24 = 0,042")
tabla(s, Cm(4.5), Cm(4.6), Cm(25), Cm(11), [
    ["Escenario", "CNN 1D", "Baseline K-S"],
    ["Test sintético (en distribución)", "1,000", "1,000"],
    [("Progresión nueva I-V-vi-IV", False, None), ("0,925", True, VERDE), "0,500"],
    [("Progresión nueva ii-V-I", False, None), ("1,000", True, VERDE), "0,500"],
    ["Timbre brillante (6 armónicos)", "0,908", "1,000"],
    ["Ruido fuerte (SNR 0 dB)", "0,808", "1,000"],
    ["Afinación 432 Hz", "0,917", "0,958"],
    [("TRANSPOSICIÓN DE 1 OCTAVA", True, ROJO), ("0,042", True, ROJO), ("1,000", True, VERDE)],
    [("Audio real (5 temas conocidos)", False, None), "0,800", ("1,000", True, VERDE)],
], ancho_cols=[Cm(14), Cm(5.5), Cm(5.5)], size=15)
caja(s, Cm(4.5), Cm(16.2), Cm(25), Cm(1.6),
     "Una canción en C mayor sigue en C mayor una octava más arriba. La etiqueta no cambió.",
     size=16, color=GRIS, align=PP_ALIGN.CENTER)

# ---------------------------------------------------------------- 9. Diagnóstico
s = nueva("Diagnóstico", "¿Qué sigue la red: la tonalidad o la frecuencia absoluta?")
imagen_centrada(s, "diagnostico_octava.png", Cm(3.8), Cm(10))
caja(s, Cm(2), Cm(14.4), AN - Cm(4), Cm(3.6),
     "Al desplazar el audio UN SEMITONO:\n"
     "predice la etiqueta transpuesta el 91,7 %  ·  acierta la correcta el 0 %",
     size=20, bold=True, color=ROJO, align=PP_ALIGN.CENTER, interlineado=1.3)

# ---------------------------------------------------------------- 10. Causa
s = nueva("La causa: shortcut learning", "No fue la arquitectura. Fue el conjunto de datos.", claro=True)
vinetas(s, Cm(1.8), Cm(4.6), Cm(30), Cm(9), [
    ("El generador ubicaba cada tonalidad SIEMPRE en la misma octava.", True),
    "Dentro del dataset, \"tonalidad\" y \"banda de frecuencias absolutas\" eran predictores equivalentes.",
    ("Ante dos soluciones equivalentes, el gradiente tomó la más fácil: detectar altura absoluta.", True),
    "Nunca necesitó construir relaciones intervalares, que es lo que define una tonalidad.",
    ("El test, generado con el mismo sesgo, no podía detectarlo.", True),
    "Una CNN 1D sobre onda cruda es equivariante en tiempo, pero NO en frecuencia.",
    "El cromagrama sí es invariante a la octava por construcción: su sesgo inductivo es el correcto.",
], size=17.5, espacio=11)
caja(s, Cm(2.5), Cm(14.6), AN - Cm(5), Cm(2.4),
     "La red aprendió a detectar altura absoluta, no tonalidad.",
     size=26, bold=True, color=ROJO, align=PP_ALIGN.CENTER)

# ---------------------------------------------------------------- 11. Conclusiones
s = nueva("Conclusiones", "Qué me llevo de este trabajo")
vinetas(s, Cm(1.8), Cm(4.5), Cm(30), Cm(11), [
    ("Con este diseño, la red NO superó a la línea de base. Esa es la respuesta honesta.", True),
    "Una métrica de test perfecta no es evidencia de aprendizaje: fue el síntoma del problema.",
    ("La línea de base no es un trámite del informe: fue el instrumento de diagnóstico.", True),
    "Sin ese empate en 100 %, el trabajo habría reportado un éxito inexistente.",
    ("Un test generado por el mismo generador que el train no mide generalización.", True),
    "Mide consistencia del generador. Detectar atajos exige evaluar fuera de distribución.",
    ("El sesgo inductivo era analizable antes de entrenar.", True),
    "El fallo estaba predicho por la arquitectura: faltaba invariancia a transposición.",
], size=17, espacio=9)

# ---------------------------------------------------------------- 12. Futuro
s = nueva("Cómo se arregla", "Las correcciones salen directamente del diagnóstico", claro=True)
tarjeta(s, Cm(1.6), Cm(4.8), Cm(15.2), Cm(5.2), "1 · Aumentación por pitch-shift",
        "Impacto alto", AZUL_CLARO, "Transponer actualizando la etiqueta rompe la correlación espuria")
tarjeta(s, Cm(17.6), Cm(4.8), Cm(15.2), Cm(5.2), "2 · Randomizar la octava",
        "2 líneas", VERDE, "Elimina el sesgo en el origen, dentro del generador")
tarjeta(s, Cm(1.6), Cm(10.8), Cm(15.2), Cm(5.2), "3 · CQT / log-mel + CNN 2D",
        "Por arquitectura", AZUL_CLARO, "Sobre CQT la transposición es traslación: la conv2D ya es equivariante")
tarjeta(s, Cm(17.6), Cm(10.8), Cm(15.2), Cm(5.2), "4 · GiantSteps completo",
        "604 temas", AZUL_CLARO, "Medición sólida sobre música real anotada")

# ---------------------------------------------------------------- 13. Cierre
s = nueva()
fondo(s, AZUL)
caja(s, Cm(2.5), Cm(5.6), AN - Cm(5), Cm(3),
     "Master of the Harmonica", size=46, bold=True, color=BLANCO, align=PP_ALIGN.CENTER)
caja(s, Cm(3), Cm(9.2), AN - Cm(6), Cm(4),
     "El sistema funciona de punta a punta.\n"
     "El aporte del trabajo no es el 100 %: es haber mostrado por qué ese 100 % no alcanzaba.",
     size=20, color=RGBColor(0xC9, 0xD9, 0xF0), align=PP_ALIGN.CENTER, interlineado=1.35)
caja(s, Cm(3), Cm(14.2), AN - Cm(6), Cm(2.5),
     "github.com/vmdanguise/master-of-the-harmonica\n\nGracias.",
     size=17, color=BLANCO, align=PP_ALIGN.CENTER, interlineado=1.4)

os.makedirs(BASE, exist_ok=True)
prs.save(SALIDA)
print("PPTX ->", SALIDA, "|", len(prs.slides.__iter__.__self__._sldIdLst), "slides")
